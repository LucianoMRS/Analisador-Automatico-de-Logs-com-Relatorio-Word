import re
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PASTA = Path(__file__).parent
TIPOS = ["INFO", "ERROR", "WARNING", "DEBUG"]
FONTE = "Times New Roman"

# Paleta de cores (hexadecimal, sem #)
COR_PRIMARIA = "1F3864"   # azul escuro: títulos e cabeçalhos
COR_ZEBRA = "EAF0F8"      # azul bem claro: linhas alternadas
COR_BORDA = "BFBFBF"      # cinza claro: bordas das tabelas
COR_TEXTO = "262626"
COR_SECUNDARIA = "595959"
COR_BRANCO = "FFFFFF"
CORES_TIPO = {
    "ERROR": "C00000",
    "WARNING": "BF8F00",
    "INFO": "2E75B6",
    "DEBUG": "7F7F7F",
}


# ----------------------------------------------------------------------
# Extração
# ----------------------------------------------------------------------
def extrair_dados(log):
    expressao = r"(\d{4}-\d{2}-\d{2})\s+(\d{2}:\d{2}:\d{2})\s+(\w+)\s+(.+)"
    resultados = re.findall(expressao, log)

    logs = []
    for data, hora, tipo, mensagem in resultados:
        logs.append({
            "Data": data,
            "Hora": hora,
            "Tipo": tipo.upper(),
            "Mensagem": mensagem.strip(),
        })

    if not logs:
        print("Não foi possível extrair os dados do log.")
    return logs


def formatar_data(data_iso):
    return datetime.strptime(data_iso, "%Y-%m-%d").strftime("%d/%m/%Y")


# ----------------------------------------------------------------------
# Funções auxiliares de estilo
# ----------------------------------------------------------------------
def aplicar_fonte(run, tamanho=12, negrito=False, cor=COR_TEXTO,
                  italico=False):
    run.font.name = FONTE
    run.font.size = Pt(tamanho)
    run.font.bold = negrito
    run.font.italic = italico
    run.font.color.rgb = RGBColor.from_string(cor)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)


def configurar_estilo_padrao(doc):
    estilo = doc.styles["Normal"]
    estilo.font.name = FONTE
    estilo.font.size = Pt(12)
    estilo.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)


def borda_inferior(paragrafo, cor, espessura=8):
    """Linha horizontal abaixo do parágrafo (usada em títulos)."""
    p_pr = paragrafo._p.get_or_add_pPr()
    borda = OxmlElement("w:pBdr")
    inferior = OxmlElement("w:bottom")
    inferior.set(qn("w:val"), "single")
    inferior.set(qn("w:sz"), str(espessura))
    inferior.set(qn("w:space"), "4")
    inferior.set(qn("w:color"), cor)
    borda.append(inferior)
    p_pr.append(borda)


def sombrear(celula, cor):
    """Cor de fundo da célula."""
    tc_pr = celula._element.get_or_add_tcPr()
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor)
    tc_pr.append(sombra)


def bordas_tabela(tabela, cor=COR_BORDA):
    tbl_pr = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for nome in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{nome}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), cor)
        bordas.append(el)
    look = tbl_pr.find(qn("w:tblLook"))
    if look is not None:
        look.addprevious(bordas)
    else:
        tbl_pr.append(bordas)


def adicionar_titulo(doc, texto):
    titulo = doc.add_heading("", level=1)
    borda_inferior(titulo, COR_PRIMARIA, 18)
    aplicar_fonte(titulo.add_run(texto), 24, True, COR_PRIMARIA)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    titulo.paragraph_format.space_before = Pt(0)
    titulo.paragraph_format.space_after = Pt(8)
    return titulo


def adicionar_subtitulo(doc, texto):
    paragrafo = doc.add_paragraph()
    aplicar_fonte(paragrafo.add_run(texto), 11, False, COR_SECUNDARIA,
                  italico=True)
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragrafo.paragraph_format.space_after = Pt(18)
    return paragrafo


def adicionar_secao(doc, texto):
    secao = doc.add_heading("", level=2)
    borda_inferior(secao, COR_PRIMARIA, 8)
    aplicar_fonte(secao.add_run(texto), 15, True, COR_PRIMARIA)
    secao.alignment = WD_ALIGN_PARAGRAPH.CENTER
    secao.paragraph_format.space_before = Pt(20)
    secao.paragraph_format.space_after = Pt(12)
    return secao


def escrever_celula(celula, texto, negrito=False, cor=COR_TEXTO,
                    alinhamento=WD_ALIGN_PARAGRAPH.CENTER):
    paragrafo = celula.paragraphs[0]
    aplicar_fonte(paragrafo.add_run(texto), 11, negrito, cor)
    paragrafo.alignment = alinhamento
    paragrafo.paragraph_format.space_before = Pt(4)
    paragrafo.paragraph_format.space_after = Pt(4)


def criar_tabela(doc, cabecalho, linhas, larguras=None, alinhamentos=None):
    """
    Cria uma tabela estilizada.
    - cabecalho: lista de títulos
    - linhas: lista de listas; cada valor pode ser "texto" ou
      ("texto", "COR_HEX", negrito) para destacar a célula
    - larguras: lista de larguras em cm (opcional)
    - alinhamentos: lista de alinhamentos por coluna (opcional)
    """
    n_cols = len(cabecalho)
    alinhamentos = alinhamentos or [WD_ALIGN_PARAGRAPH.CENTER] * n_cols

    tabela = doc.add_table(rows=1, cols=n_cols)
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabela.autofit = False
    bordas_tabela(tabela)

    for i, titulo in enumerate(cabecalho):
        celula = tabela.rows[0].cells[i]
        sombrear(celula, COR_PRIMARIA)
        escrever_celula(celula, titulo, True, COR_BRANCO,
                        WD_ALIGN_PARAGRAPH.CENTER)

    for n, linha in enumerate(linhas):
        celulas = tabela.add_row().cells
        for i, valor in enumerate(linha):
            if isinstance(valor, tuple):
                texto, cor, negrito = valor
            else:
                texto, cor, negrito = valor, COR_TEXTO, False
            if n % 2 == 1:
                sombrear(celulas[i], COR_ZEBRA)
            escrever_celula(celulas[i], texto, negrito, cor, alinhamentos[i])

    if larguras:
        for linha in tabela.rows:
            for i, largura in enumerate(larguras):
                linha.cells[i].width = Cm(largura)

    return tabela


def adicionar_rodape(doc):
    paragrafo = doc.sections[0].footer.paragraphs[0]
    paragrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    aplicar_fonte(paragrafo.add_run("Relatório de Análise de Logs  |  Página "),
                  10, False, COR_SECUNDARIA)

    run = paragrafo.add_run()
    aplicar_fonte(run, 10, False, COR_SECUNDARIA)
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instrucao = OxmlElement("w:instrText")
    instrucao.set(qn("xml:space"), "preserve")
    instrucao.text = " PAGE "
    fim = OxmlElement("w:fldChar")
    fim.set(qn("w:fldCharType"), "end")
    run._r.append(inicio)
    run._r.append(instrucao)
    run._r.append(fim)


# ----------------------------------------------------------------------
# Relatório
# ----------------------------------------------------------------------
def criar_relatorio():
    caminho_log = PASTA / "logs.txt"

    if not caminho_log.exists():
        print(f"Arquivo não encontrado: {caminho_log}")
        print("Coloque o logs.txt na mesma pasta do main.py.")
        return

    with open(caminho_log, encoding="utf-8") as arquivo:
        logs = extrair_dados(arquivo.read())

    if not logs:
        return

    doc = Document()
    configurar_estilo_padrao(doc)
    adicionar_rodape(doc)

    # Contagens
    contagem = {}
    for log in logs:
        contagem[log["Tipo"]] = contagem.get(log["Tipo"], 0) + 1
    total = len(logs)

    datas = sorted({log["Data"] for log in logs})
    agora = datetime.now()

    # Capa / título
    adicionar_titulo(doc, "Relatório de Análise de Logs")
    adicionar_subtitulo(
        doc,
        f"Gerado em {agora.strftime('%d/%m/%Y')} às {agora.strftime('%H:%M')}"
        f"  •  {total} registros analisados"
        f"  •  Período: {formatar_data(datas[0])} a {formatar_data(datas[-1])}",
    )

    # Resumo por tipo
    adicionar_secao(doc, "Total de ocorrências por tipo")
    tipos_ordenados = [t for t in TIPOS if t in contagem]
    tipos_ordenados += sorted(t for t in contagem if t not in TIPOS)

    linhas_resumo = []
    for tipo in tipos_ordenados:
        qtd = contagem[tipo]
        cor = CORES_TIPO.get(tipo, COR_TEXTO)
        linhas_resumo.append([
            (tipo, cor, True),
            str(qtd),
            f"{qtd / total * 100:.1f}%",
        ])
    criar_tabela(doc, ["Tipo", "Quantidade", "Percentual"], linhas_resumo,
                 larguras=[5.3, 5.3, 5.3])

    # Logs de erro
    erros = [log for log in logs if log["Tipo"] == "ERROR"]
    if erros:
        adicionar_secao(doc, "Logs de erro extraídos")
        linhas_erros = [
            [formatar_data(log["Data"]), log["Hora"], log["Mensagem"]]
            for log in erros
        ]
        criar_tabela(
            doc,
            ["Data", "Hora", "Mensagem"],
            linhas_erros,
            larguras=[2.8, 2.2, 10.9],
            alinhamentos=[
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.LEFT,
            ],
        )

    doc.add_page_break()

    # Registros por dia
    registro_por_dia = {}
    for log in logs:
        dia = registro_por_dia.setdefault(log["Data"], {})
        dia[log["Tipo"]] = dia.get(log["Tipo"], 0) + 1

    adicionar_secao(doc, "Registros por dia")
    linhas_dias = []
    for data in sorted(registro_por_dia):
        dia = registro_por_dia[data]
        linha = [(formatar_data(data), COR_TEXTO, True)]
        for tipo in TIPOS:
            qtd = dia.get(tipo, 0)
            if tipo == "ERROR" and qtd > 0:
                linha.append((str(qtd), CORES_TIPO["ERROR"], True))
            else:
                linha.append(str(qtd))
        linha.append((str(sum(dia.values())), COR_PRIMARIA, True))
        linhas_dias.append(linha)

    criar_tabela(doc, ["Data"] + TIPOS + ["Total"], linhas_dias,
                 larguras=[3.4, 2.1, 2.1, 2.5, 2.1, 2.1])

    # Salvar ao lado do script
    nome_arquivo = f"relatorio_logs_{agora.strftime('%Y-%m-%d')}.docx"
    caminho_saida = PASTA / nome_arquivo
    doc.save(caminho_saida)
    print(f"Relatório gerado: {caminho_saida}")


if __name__ == "__main__":
    criar_relatorio()